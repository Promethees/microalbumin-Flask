"""Tests for the native calibration chart in /export_report_excel.

The chart is a category-axis LineChart (markers only) whose X labels are a
"nice numbers" tick series (0, 100, 200, …) over the concentration range. These
tests guard:
  * string standards points coerced to real numbers (plottable, locale-aware);
  * the axis labelled with round, evenly-spaced ticks (not the raw helper Xs);
  * two distinct markers, no connecting lines, no native trendline.
"""
import io
import os
import re
import sys
import zipfile

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

from main import app  # noqa: E402
from routes.file_routes import _nice_axis_ticks  # noqa: E402

# Helper layout: AA(27)=conc (nice-tick label), AB(28)=std_y, AC(29)=fit_y.
CAT, STD_Y, FIT_Y = 27, 28, 29


@pytest.fixture
def client():
    app.config['TESTING'] = True
    with app.test_client() as c:
        yield c


def _helper_sheet(wb):
    for ws in wb.worksheets:
        if ws.cell(1, CAT).value == 'conc':
            return ws
    return None


def _read_helper_rows(ws):
    """Read (conc, std_y, fit_y) rows until a fully-empty row ends the data."""
    rows, r = [], 2
    while any(ws.cell(r, c).value is not None for c in (CAT, STD_Y, FIT_Y)):
        rows.append((ws.cell(r, CAT).value,
                     ws.cell(r, STD_Y).value,
                     ws.cell(r, FIT_Y).value))
        r += 1
    return rows


def _fit_grid(lo, hi, n=24):
    """A uniform fit sampling (smooth-curve stand-in) over [lo, hi]."""
    step = (hi - lo) / (n - 1)
    return [{'x': lo + i * step, 'y': 0.1 * ((lo + i * step) / (300 + lo + i * step))}
            for i in range(n)]


def _calibrate_item(points, fit):
    return {
        'filename': 'cal.csv',
        'mode': 'calibrate',
        'csv_columns': ['Concentration', 'maxRate'],
        'csv_rows': [{'Concentration': '5', 'maxRate': '0.00738'}],
        'chart_series': [{
            'label': 'maxRate — Michaelis-Menten',
            'title': 'maxRate — Michaelis-Menten',
            'algo': 'Michaelis-Menten',
            'xLabel': 'Concentration (ng/µL)', 'yLabel': 'maxRate',
            'points': points,
            'fit': fit,
        }],
    }


# ── Nice-numbers tick algorithm ──────────────────────────────────────────────

def test_nice_axis_ticks_round_and_evenly_spaced():
    assert _nice_axis_ticks(5, 500) == [0, 100, 200, 300, 400, 500]
    assert _nice_axis_ticks(5, 50) == [0, 10, 20, 30, 40, 50]
    assert _nice_axis_ticks(100, 560) == [100, 200, 300, 400, 500, 600]
    ticks = _nice_axis_ticks(0.2, 2.1)
    assert ticks[0] == 0.0 and ticks[-1] >= 2.1
    steps = {round(b - a, 10) for a, b in zip(ticks, ticks[1:])}
    assert len(steps) == 1, f"ticks not evenly spaced: {ticks}"


# ── Cell writing ─────────────────────────────────────────────────────────────

def test_string_points_written_as_numbers(client):
    # The standards arrive as strings (as they do from the CSV-parsed client).
    item = _calibrate_item(
        points=[{'x': '5', 'y': '0.00738'}, {'x': '50', 'y': '0.01884'},
                {'x': '500', 'y': '0.085'}],
        fit=_fit_grid(0, 500),
    )
    rv = client.post('/export_report_excel', json={'items': [item], 'subject': 'T'})
    assert rv.status_code == 200

    ws = _helper_sheet(load_workbook(io.BytesIO(rv.data)))
    assert ws is not None, "category helper header not found"
    rows = _read_helper_rows(ws)

    # Standards survive as real numbers in std_y (the string→float coercion).
    std = sorted(ys for (c, ys, yf) in rows if ys is not None)
    assert std == pytest.approx([0.00738, 0.01884, 0.085])
    assert all(isinstance(v, (int, float)) for v in std)
    # Fit points survive as numbers too.
    assert all(isinstance(yf, (int, float)) for (c, ys, yf) in rows if yf is not None)
    # The category labels are round nice-numbers ticks, not the raw standard Xs.
    cats = sorted(c for (c, ys, yf) in rows if c is not None)
    assert cats == [0.0, 100.0, 200.0, 300.0, 400.0, 500.0]


def test_non_numeric_points_are_skipped(client):
    item = _calibrate_item(
        points=[{'x': '5', 'y': '0.1'}, {'x': 'NONE', 'y': '0.2'}],
        fit=[],
    )
    rv = client.post('/export_report_excel', json={'items': [item]})
    assert rv.status_code == 200

    ws = _helper_sheet(load_workbook(io.BytesIO(rv.data)))
    rows = _read_helper_rows(ws)
    # Only the one valid standard's Y is written; the 'NONE' point is dropped.
    assert [ys for (c, ys, yf) in rows if ys is not None] == [pytest.approx(0.1)]
    assert len(rows) == 1


# ── Chart structure ──────────────────────────────────────────────────────────

def _chart_xml(xlsx_bytes):
    with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as z:
        names = sorted(n for n in z.namelist() if n.startswith('xl/charts/chart'))
        assert names, "no chart embedded in the workbook"
        return z.read(names[0]).decode('utf-8')


def test_category_line_chart_with_nice_tick_labels(client):
    item = _calibrate_item(
        points=[{'x': '5', 'y': '0.00738'}, {'x': '50', 'y': '0.01884'},
                {'x': '500', 'y': '0.085'}],
        fit=_fit_grid(0, 500),
    )
    rv = client.post('/export_report_excel', json={'items': [item]})
    assert rv.status_code == 200
    xml = _chart_xml(rv.data)

    # A category-axis LINE chart (not a scatter) — that's what lets us relabel X.
    assert '<lineChart>' in xml and '<scatterChart>' not in xml
    assert xml.count('<catAx>') == 1 and xml.count('<valAx>') == 1
    assert xml.count('<ser>') == 2          # standards + fit
    assert xml.count('<cat>') == 2 and xml.count('<val>') == 2
    assert xml.count('<numCache>') == 4     # cat + val × 2 series
    # The category labels are the round nice-numbers ticks.
    cat_blocks = re.findall(r'<cat>.*?</cat>', xml, re.S)
    cat_vals = sorted(float(v) for v in re.findall(r'<pt idx="\d+"><v>([^<]+)</v>', cat_blocks[0]))
    assert cat_vals == [0.0, 100.0, 200.0, 300.0, 400.0, 500.0]
    # Force every slot's label (Sheets self-skips otherwise) and keep them flat.
    assert '<tickLblSkip val="1"' in xml and 'rot="0"' in xml
    # Two DISTINCT marker shapes, neither with a connecting line.
    assert '<symbol val="none"' not in xml
    assert xml.count('<symbol val="diamond"') == 1  # standards
    assert xml.count('<symbol val="dot"') == 1      # fit — smallest marker
    assert xml.count('<a:noFill') == 2, "a series carries a connecting line"
    assert 'trendline' not in xml.lower()


def test_no_native_trendline_for_michaelis_menten(client):
    item = _calibrate_item(
        points=[{'x': '5', 'y': '0.00738'}, {'x': '50', 'y': '0.01884'}],
        fit=_fit_grid(0, 200, 12),
    )  # _calibrate_item already uses algo 'Michaelis-Menten'
    rv = client.post('/export_report_excel', json={'items': [item]})
    assert 'trendline' not in _chart_xml(rv.data).lower()
